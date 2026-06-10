"""
WAF 测试报告生成脚本 - 将截图插入 Word 模板
用法: python3 report_skill.py --screenshots-dir <dir> [--template-path <path>]
"""
import argparse
import json
import os
import sys
from datetime import datetime


# 测试类型 → 文件名关键词 + 表格内容关键词（从 report_views.py 提取）
TEST_PATTERNS = {
    'SQL注入攻击': {
        'file_keywords': ['sql_event'],
        'table_keywords': ['sql注入', 'sql injection', 'union select', 'or 1=1', 'sql'],
    },
    'XSS攻击': {
        'file_keywords': ['xss_event'],
        'table_keywords': ['xss', 'cross site scripting', '<script>', 'javascript'],
    },
    'CSRF攻击': {
        'file_keywords': ['csrf_event'],
        'table_keywords': ['csrf', 'cross site request forgery', '跨站请求伪造'],
    },
    'SSRF攻击': {
        'file_keywords': ['ssrf_event'],
        'table_keywords': ['ssrf', 'server side request forgery', '服务端请求伪造', '服务器端请求伪造'],
    },
    'ASP代码注入攻击': {
        'file_keywords': ['asp_event'],
        'table_keywords': ['asp', 'asp测试', 'asp攻击', 'asp代码'],
    },
    '命令注入攻击': {
        'file_keywords': ['code_event'],
        'table_keywords': ['命令注入', 'command injection', '命令执行', '代码注入', 'code injection', '执行任意命令', '代码执行', 'code测试'],
    },
    'Java反序列化攻击': {
        'file_keywords': ['java_event'],
        'table_keywords': ['Java反序列化', 'java deserialization', 'Java反序列化攻击检测'],
    },
    'Java代码注入攻击': {
        'file_keywords': ['java_code_event'],
        'table_keywords': ['Java代码注入', 'Java code injection', 'Java代码注入攻击检测'],
    },
    'PHP代码注入攻击': {
        'file_keywords': ['php_event'],
        'table_keywords': ['PHP代码注入', 'PHP code injection', 'PHP代码注入攻击检测', 'PHP注入'],
    },
    'PHP反序列化攻击': {
        'file_keywords': ['php_deserialization_event'],
        'table_keywords': ['PHP反序列化', 'PHP deserialization', 'PHP反序列化攻击检测', 'PHP反序列化对象解析'],
    },
    '文件包含攻击': {
        'file_keywords': ['fileinclude_event', 'file_event'],
        'table_keywords': ['文件包含', 'file inclusion', 'path traversal', '../', '文件包含攻击'],
    },
    '文件上传攻击': {
        'file_keywords': ['fileupload_event', 'upload_event'],
        'table_keywords': ['文件上传', 'file upload', '上传测试', '上传攻击'],
    },
    '机器人检测模块': {
        'file_keywords': ['robot_event'],
        'table_keywords': ['机器人检测模块', 'bot检测算法'],
    },
    '情报模块测试': {
        'file_keywords': ['intelligence_module_event'],
        'table_keywords': ['测试文件', '备份文件', '代码仓库', '敏感文件'],
    },
    '服务器响应测试': {
        'file_keywords': ['server_event'],
        'table_keywords': ['服务器响应', 'server response', '服务器检测', '响应测试'],
    },
    '机器人响应测试': {
        'file_keywords': ['robot_event'],
        'table_keywords': ['机器人响应', 'robot response', '机器人检测', '爬虫检测'],
    },
    'Base64测试': {
        'file_keywords': ['base64_event'],
        'table_keywords': ['base 64', '成功解码', 'WAF的防护压力'],
    },
    'URL编码测试': {
        'file_keywords': ['url_event'],
        'table_keywords': ['URL编码'],
    },
    'JSON解析测试': {
        'file_keywords': ['json_event'],
        'table_keywords': ['JSON编码'],
    },
    '16进制编码': {
        'file_keywords': ['onesix_event'],
        'table_keywords': ['HEX编码'],
    },
    '斜杠反转译': {
        'file_keywords': ['backslash_event'],
        'table_keywords': ['Unicode Escape编码', 'Unicode Escape解码'],
    },
    'XML编码测试': {
        'file_keywords': ['xml_event'],
        'table_keywords': ['XML的请求', 'XML解析', '成功解析'],
    },
    'UTF-7编码': {
        'file_keywords': ['utf7_event'],
        'table_keywords': ['UTF-7编码'],
    },
    '多层解码测试': {
        'file_keywords': ['more_value_event'],
        'table_keywords': ['多层编码', '成功解码', '多层编码解析'],
    },
    'URL匹配': {
        'file_keywords': ['custom_urlpath_event'],
        'table_keywords': ['自定义URL匹配', 'URL规则'],
    },
    'Query匹配': {
        'file_keywords': ['custom_decoded_query_event'],
        'table_keywords': ['自定义Query匹配规则', 'Query规则'],
    },
    '路径匹配': {
        'file_keywords': ['custom_decoded_path_event'],
        'table_keywords': ['自定义路径匹配规则', '路径规则'],
    },
    'Body匹配': {
        'file_keywords': ['custom_request_body_event'],
        'table_keywords': ['自定义Body匹配规则', 'Body规则'],
    },
    'Origin匹配': {
        'file_keywords': ['custom_origin_event'],
        'table_keywords': ['自定义Origin匹配规则', 'Origin规则'],
    },
    'X-Forwarded-For匹配': {
        'file_keywords': ['custom_x_forwarded_for_event'],
        'table_keywords': ['自定义X-Forwarded-For匹配规则', 'X-Forwarded-For规则'],
    },
    'Header匹配': {
        'file_keywords': ['custom_request_header_event'],
        'table_keywords': ['自定义Header匹配规则', 'Header规则'],
    },
    'Content-Type匹配': {
        'file_keywords': ['custom_content_type_event'],
        'table_keywords': ['自定义Content-Type匹配规则', 'Content-Type规则来'],
    },
    '请求方法匹配': {
        'file_keywords': ['custom_custom_header_event'],
        'table_keywords': ['自定义请求方法匹配规则', '匹配请求方法规则'],
    },
    'User-Agent匹配': {
        'file_keywords': ['custom_user_agent_event'],
        'table_keywords': ['自定义User-Agent匹配规则', 'User-Agent规则'],
    },
    'Cookie匹配': {
        'file_keywords': ['custom_cookie_raw_event'],
        'table_keywords': ['自定义Cookie匹配规则'],
    },
    'Referrer匹配': {
        'file_keywords': ['custom_referer_event'],
        'table_keywords': ['自定义Referrer匹配规则', 'Referrer规则'],
    },
    'WAF登录测试': {
        'file_keywords': ['waf_login'],
        'table_keywords': ['waf', '登录', 'login', '防护'],
    },
}


def classify_screenshots(screenshot_paths):
    """按文件名关键词将截图分类到对应测试类型"""
    # 按关键词长度降序排列，确保更具体的优先匹配
    sorted_patterns = []
    for test_type, config in TEST_PATTERNS.items():
        for keyword in config['file_keywords']:
            sorted_patterns.append((test_type, keyword, len(keyword)))
    sorted_patterns.sort(key=lambda x: x[2], reverse=True)

    classified = {}
    for path in screenshot_paths:
        filename = os.path.basename(path).lower()
        matched = False
        for test_type, keyword, _ in sorted_patterns:
            if keyword.lower() in filename:
                classified.setdefault(test_type, []).append(path)
                print(f"  截图分类: {os.path.basename(path)} -> {test_type}")
                matched = True
                break
        if not matched:
            classified.setdefault('other', []).append(path)
            print(f"  截图分类: {os.path.basename(path)} -> 未分类")
    return classified


def match_tables_to_screenshots(doc, classified):
    """扫描 Word 模板中的测试表格，匹配对应截图"""
    target_tables = []
    matched_types = []

    for i, table in enumerate(doc.tables):
        table_text = ""
        for row in table.rows:
            for cell in row.cells:
                table_text += cell.text + " "

        indicators = ['测试项目', '预期结果', '执行步骤', '判定结果', '测试人员']
        if not any(ind in table_text for ind in indicators):
            continue

        # 按匹配度打分
        text_lower = table_text.lower()
        scores = {}
        for test_type, config in TEST_PATTERNS.items():
            score = sum(len(kw) for kw in config['table_keywords'] if kw.lower() in text_lower)
            if score > 0:
                scores[test_type] = score

        if not scores:
            continue

        test_type = max(scores.items(), key=lambda x: x[1])[0]
        if test_type in classified and classified[test_type]:
            target_tables.append({
                'table': table,
                'table_index': i,
                'test_type': test_type,
                'screenshots': classified[test_type],
            })
            if test_type not in matched_types:
                matched_types.append(test_type)
            print(f"  表格匹配: 表格 {i + 1} -> {test_type} (有截图)")

    return target_tables, matched_types


def insert_screenshots_into_doc(target_tables):
    """在匹配到的表格「预期结果」行中插入截图（basic + detail 各一张）"""
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    inserted = 0
    for info in target_tables:
        table = info['table']
        screenshots = info['screenshots']
        if not screenshots:
            continue

        # 按 type 排序：basic 在前，detail 在后
        screenshots_sorted = sorted(screenshots, key=lambda s: 0 if '_basic_' in s else 1)

        # 找到「预期结果」行
        result_row_idx = -1
        for row_idx, row in enumerate(table.rows):
            row_text = ''.join(cell.text for cell in row.cells)
            if '预期结果' in row_text:
                result_row_idx = row_idx
                break

        if result_row_idx == -1 and len(table.rows) >= 4:
            result_row_idx = len(table.rows) - 2

        if result_row_idx == -1:
            print(f"  表格 {info['table_index'] + 1}: 未找到预期结果行")
            continue

        result_row = table.rows[result_row_idx]
        target_col = min(1, len(result_row.cells) - 1)
        target_cell = result_row.cells[target_col]

        try:
            paragraph = target_cell.paragraphs[0] if target_cell.paragraphs else target_cell.add_paragraph()
            if paragraph.text.strip():
                paragraph.add_run().add_break()
                paragraph.add_run().add_break()

            # 插入所有截图（basic + detail）
            for i, screenshot_path in enumerate(screenshots_sorted):
                if i > 0:
                    # 在两张图之间加分隔
                    sep_run = paragraph.add_run()
                    sep_run.add_break()
                    sep_run.add_break()

                run = paragraph.add_run()
                run.add_picture(screenshot_path, width=Inches(4.5))
                inserted += 1
                label = "基本信息" if "_basic_" in screenshot_path else "详细详情"
                print(f"  表格 {info['table_index'] + 1} ({info['test_type']}): 插入{label}截图 {os.path.basename(screenshot_path)}")

            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception as e:
            print(f"  表格 {info['table_index'] + 1} 插入失败: {e}")

    return inserted


def generate_report(screenshots_dir, template_path=None, output_path=None):
    """主入口：生成 Word 测试报告"""
    from docx import Document

    # 默认模板路径
    if not template_path:
        script_dir = os.path.dirname(os.path.abspath(__file__))
        template_path = os.path.join(script_dir, '..', 'doc', 'H20-单机反代测试方案v1.2.docx')
        template_path = os.path.normpath(template_path)

    if not os.path.exists(template_path):
        print(f"错误: 模板文件不存在: {template_path}")
        sys.exit(1)

    # 收集所有截图（basic + detail）
    all_screenshots = []
    for f in os.listdir(screenshots_dir):
        if f.lower().endswith('.png'):
            all_screenshots.append(os.path.join(screenshots_dir, f))

    if not all_screenshots:
        print(f"错误: 目录中没有截图文件: {screenshots_dir}")
        sys.exit(1)

    print(f"\n找到 {len(all_screenshots)} 张截图")

    # 分类
    print("截图分类:")
    classified = classify_screenshots(all_screenshots)
    for t, paths in classified.items():
        if paths:
            print(f"  {t}: {len(paths)} 张")

    # 加载模板并匹配
    doc = Document(template_path)
    print("\n表格匹配:")
    target_tables, matched_types = match_tables_to_screenshots(doc, classified)
    print(f"\n匹配到 {len(target_tables)} 个表格")

    # 插入截图
    print("插入截图:")
    inserted = insert_screenshots_into_doc(target_tables)

    # 保存
    if not output_path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = os.path.join(screenshots_dir, f"测试报告_{timestamp}.docx")

    doc.save(output_path)
    print(f"\n报告已生成: {output_path}")
    print(f"共插入 {inserted} 张截图到 {len(target_tables)} 个表格")
    return output_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description='WAF 测试报告生成工具')
    parser.add_argument('--screenshots-dir', required=True, help='截图所在目录')
    parser.add_argument('--template-path', default=None, help='Word 模板路径（默认自动定位）')
    parser.add_argument('--output-path', default=None, help='输出报告路径（默认保存到截图目录）')
    args = parser.parse_args()

    if not os.path.isdir(args.screenshots_dir):
        print(f"错误: 截图目录不存在: {args.screenshots_dir}")
        sys.exit(1)

    generate_report(args.screenshots_dir, args.template_path, args.output_path)
